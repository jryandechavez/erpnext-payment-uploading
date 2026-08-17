import csv
import io
from collections import OrderedDict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

import frappe
from frappe import _
from frappe.utils import flt, getdate, nowdate


@frappe.whitelist()
def preview(file_url):
    _require_create_permission()
    rows = _read_rows(file_url)
    if not rows:
        frappe.throw(_("The uploaded file has no payment rows."))

    names = list({row.invoice_no for row in rows if row.invoice_no})
    invoices = {
        row.name: row
        for row in frappe.get_all(
            "Sales Invoice",
            filters={"name": ["in", names]},
            fields=["name", "customer", "customer_name", "outstanding_amount", "currency", "docstatus", "debit_to"],
            limit_page_length=0,
        )
    }
    total_ewt = _money(sum((Decimal(str(row.ewt_amount)) for row in rows), Decimal("0")))
    allocated_ewt = Decimal("0")
    result = []
    for index, row in enumerate(rows):
        row_ewt = total_ewt - allocated_ewt if index == len(rows) - 1 else _money(row.ewt_amount)
        allocated_ewt += row_ewt
        invoice = invoices.get(row.invoice_no)
        errors, warnings = [], []
        if not row.cheque_no:
            errors.append(_("Check # is required"))
        if not row.invoice_no:
            errors.append(_("Invoice # is required"))
        elif not invoice:
            errors.append(_("Sales Invoice was not found"))
        elif invoice.docstatus != 1:
            errors.append(_("Sales Invoice is not submitted"))
        elif flt(invoice.outstanding_amount) <= 0:
            errors.append(_("Sales Invoice has no outstanding balance"))
        if abs(row.basis_amount - row.ewt_amount - row.cheque_amount) > Decimal("0.02"):
            errors.append(_("Calculated net amount does not equal basis less EWT"))
        if invoice and abs(row.invoice_amount - Decimal(str(invoice.outstanding_amount or 0))) > Decimal("0.02"):
            warnings.append(_("Uploaded amount differs from current outstanding"))

        result.append(
            {
                **row,
                "invoice_amount": float(row.invoice_amount),
                "basis_amount": float(row.basis_amount),
                "ewt_amount": float(row_ewt),
                "cheque_amount": float(row.cheque_amount),
                "difference_amount": float(row.difference_amount),
                "customer": invoice.customer if invoice else None,
                "customer_name": invoice.customer_name if invoice else None,
                "currency": invoice.currency if invoice else None,
                "outstanding": flt(invoice.outstanding_amount) if invoice else 0,
                "write_off_amount": float(
                    _money(invoice.outstanding_amount) - _money(row.invoice_amount) - row_ewt
                )
                if invoice
                else 0,
                "status": "Invalid" if errors else ("Warning" if warnings else "Valid"),
                "message": "; ".join(errors + warnings),
            }
        )
    result.sort(key=lambda row: (row.get("customer") or "", row.get("cheque_no") or "", row.get("row_no") or 0))
    return {"rows": result, "invalid_rows": sum(r["status"] == "Invalid" for r in result)}


@frappe.whitelist()
def create_journal_entries(
    rows,
    debits,
    company,
    posting_date,
    check_amount_received,
    ewt_account=None,
    write_off_account=None,
):
    _require_create_permission()
    rows = frappe.parse_json(rows) if isinstance(rows, str) else rows
    debits = frappe.parse_json(debits) if isinstance(debits, str) else debits
    if not rows or not debits or not company or not posting_date or not check_amount_received:
        frappe.throw(_("Rows, debit entries, Company, Posting Date, and Check Amount Received are required."))

    input_rows = [frappe._dict(value) for value in rows]
    uploaded_paid_total = _money(sum((Decimal(str(row.invoice_amount)) for row in input_rows), Decimal("0")))
    if uploaded_paid_total != _money(check_amount_received):
        frappe.throw(
            _("Check/Bank/Debit Amount ({0}) must equal uploaded Paid Amount total ({1}).").format(
                _money(check_amount_received), uploaded_paid_total
            )
        )
    invoice_names = list({row.invoice_no for row in input_rows})
    all_invoices = {
        row.name: row
        for row in frappe.get_all(
            "Sales Invoice",
            filters={"name": ["in", invoice_names], "docstatus": 1, "company": company},
            fields=["name", "customer", "outstanding_amount", "currency", "debit_to", "due_date"],
            limit_page_length=0,
        )
    }
    if len(all_invoices) != len(invoice_names):
        frappe.throw(_("One or more invoices are missing, not submitted, or belong to another company."))

    company_currency = frappe.get_cached_value("Company", company, "default_currency")
    invoice_currencies = {invoice.currency for invoice in all_invoices.values()}
    if invoice_currencies != {company_currency}:
        frappe.throw(
            _("This uploader currently supports company-currency invoices only ({0}). Found: {1}").format(
                company_currency, ", ".join(sorted(invoice_currencies))
            )
        )

    account_names = {row.get("account") for row in debits if row.get("account")}
    account_names.update(invoice.debit_to for invoice in all_invoices.values())
    if ewt_account:
        account_names.add(ewt_account)
    if write_off_account:
        account_names.add(write_off_account)
    account_currencies = {
        row.name: row.account_currency or company_currency
        for row in frappe.get_all(
            "Account",
            filters={"name": ["in", list(account_names)], "company": company, "is_group": 0},
            fields=["name", "account_currency"],
            limit_page_length=0,
        )
    }
    if len(account_currencies) != len(account_names):
        frappe.throw(_("One or more selected accounts are invalid for company {0}.").format(company))
    foreign_accounts = [name for name, currency in account_currencies.items() if currency != company_currency]
    if foreign_accounts:
        frappe.throw(
            _("This uploader currently supports company-currency accounts only. Check: {0}").format(
                ", ".join(foreign_accounts)
            )
        )

    # One uploaded workbook creates one Journal Entry. Customer Ref No remains
    # an invoice-level audit value in User Remark, not a grouping key.
    input_rows.sort(key=lambda row: row.row_no or 0)
    credit_by_invoice = OrderedDict()
    for row in input_rows:
        if row.invoice_no not in credit_by_invoice:
            invoice = all_invoices[row.invoice_no]
            outstanding = _money(invoice.outstanding_amount)
            if outstanding <= 0:
                frappe.throw(_("Invoice {0} has no outstanding balance.").format(row.invoice_no))
            credit_by_invoice[row.invoice_no] = {
                "outstanding": outstanding,
                "paid": Decimal("0"),
                "ewt_raw": Decimal("0"),
            }
        credit_by_invoice[row.invoice_no]["paid"] += Decimal(str(row.invoice_amount))
        credit_by_invoice[row.invoice_no]["ewt_raw"] += Decimal(str(row.ewt_amount))

    debit_rows = [frappe._dict(row) for row in debits]
    for debit in debit_rows:
        if not debit.account or _money(debit.amount) <= 0:
            frappe.throw(_("Every debit row requires an account and positive amount."))
    paid_debit_total = sum((_money(row.amount) for row in debit_rows), Decimal("0"))
    ewt_total = _money(sum((Decimal(str(row.ewt_amount)) for row in input_rows), Decimal("0")))
    if ewt_total and not ewt_account:
        frappe.throw(_("Select an EWT Account."))

    # Required order: configured debit(s), invoice references, write-off, EWT.
    accounts = []
    for debit in debit_rows:
        account_row = {
            "account": debit.account,
            "account_currency": company_currency,
            "exchange_rate": 1,
            "debit_in_account_currency": float(_money(debit.amount)),
            "user_remark": debit.get("remark") or _("Check/Bank/Debit Amount"),
        }
        if debit.get("party_type") and debit.get("party"):
            account_row.update({"party_type": debit.party_type, "party": debit.party})
        accounts.append(account_row)

    invoice_names_in_order = list(credit_by_invoice)
    allocated_ewt = Decimal("0")
    for index, name in enumerate(invoice_names_in_order):
        values = credit_by_invoice[name]
        values["ewt"] = (
            ewt_total - allocated_ewt
            if index == len(invoice_names_in_order) - 1
            else _money(values["ewt_raw"])
        )
        allocated_ewt += values["ewt"]

    for name, values in credit_by_invoice.items():
        invoice = all_invoices[name]
        accounts.append(
            {
                "account": invoice.debit_to,
                "account_currency": company_currency,
                "exchange_rate": 1,
                "party_type": "Customer",
                "party": invoice.customer,
                "credit_in_account_currency": float(values["outstanding"]),
                "reference_type": "Sales Invoice",
                "reference_name": name,
                "reference_due_date": invoice.due_date,
                "user_remark": name,
            }
        )

    # Write-off is calculated and tagged per invoice. Multiple invoice
    # differences therefore produce multiple Journal Entry Account rows.
    for name, values in credit_by_invoice.items():
        invoice = all_invoices[name]
        difference = _money(values["paid"]) + values["ewt"] - values["outstanding"]
        if abs(difference) <= Decimal("0.005"):
            continue
        write_off_row = {
            "account": write_off_account,
            "account_currency": company_currency,
            "exchange_rate": 1,
            "party_type": "Customer",
            "party": invoice.customer,
            "reference_type": "Sales Invoice",
            "reference_name": name,
            "reference_due_date": invoice.due_date,
            "user_remark": _("Write-off for {0}").format(name),
        }
        if difference > 0:
            write_off_row["credit_in_account_currency"] = float(difference)
        else:
            write_off_row["debit_in_account_currency"] = float(abs(difference))
        accounts.append(write_off_row)

    if ewt_total:
        accounts.append(
            {
                "account": ewt_account,
                "account_currency": company_currency,
                "exchange_rate": 1,
                "debit_in_account_currency": float(ewt_total),
                "user_remark": _("EWT"),
            }
        )

    customer_refs = list(dict.fromkeys(_clean_id(row.cheque_no) for row in input_rows))
    first = input_rows[0]
    remark_lines = [
            "\t".join(
                [
                    _clean_id(row.cheque_no),
                    _display_date(posting_date),
                    _display_date(row.cheque_date),
                    row.invoice_no,
                    "",
                    "{:,.2f}".format(Decimal(str(row.invoice_amount))),
                ]
            )
            for row in input_rows
        ]
    doc = frappe.get_doc(
            {
                "doctype": "Journal Entry",
                "voucher_type": "Bank Entry",
                "company": company,
                "posting_date": getdate(posting_date),
                "cheque_no": ", ".join(customer_refs)[:140],
                "cheque_date": getdate(first.cheque_date or nowdate()),
                "user_remark": "\n".join(remark_lines),
                "accounts": accounts,
            }
        )
    doc.insert()
    return {"journal_entries": [doc.name], "count": 1}


def _read_rows(file_url):
    file_doc = frappe.get_doc("File", {"file_url": file_url})
    content = file_doc.get_content()
    filename = (file_doc.file_name or file_url).lower()
    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig") if isinstance(content, bytes) else content
        raw = list(csv.DictReader(io.StringIO(text)))
    elif filename.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if "Sheet1" in workbook.sheetnames:
            sheet1 = workbook["Sheet1"]
            if _header(sheet1.cell(row=3, column=4).value) in {"si", "invoice", "invoice #"}:
                return _read_sheet1_columns(sheet1)
        sheet = workbook["july 24 2026"] if "july 24 2026" in workbook.sheetnames else workbook.active
        values = list(sheet.iter_rows(values_only=True))
        header_index = next((i for i, values_row in enumerate(values) if "invoice #" in {_header(v) for v in values_row}), None)
        if header_index is None:
            frappe.throw(_("Could not find a row containing the Invoice # header."))
        headers = [str(v or "") for v in values[header_index]]
        raw = [dict(zip(headers, values_row)) for values_row in values[header_index + 1 :] if any(v not in (None, "") for v in values_row)]
    else:
        frappe.throw(_("Please upload a CSV or XLSX file."))

    rows = []
    for index, source in enumerate(raw, start=2):
        value = {_header(k): v for k, v in source.items()}
        try:
            invoice_amount = _decimal(value.get("invoice amount"))
            ewt_amount = _decimal(value.get("ewt1%") or value.get("ewt 1%"))
            cheque_amount = _decimal(value.get("check amount") or value.get("cheque amount"))
        except InvalidOperation:
            frappe.throw(_("Row {0} contains an invalid monetary amount.").format(index))
        rows.append(
            frappe._dict(
                row_no=index,
                cheque_no=_clean_id(value.get("check #") or value.get("cheque #")),
                cheque_date=str(getdate(value.get("check date") or value.get("cheque date"))) if value.get("check date") or value.get("cheque date") else None,
                invoice_no=_clean_id(value.get("invoice #")),
                invoice_amount=invoice_amount,
                basis_amount=invoice_amount,
                ewt_amount=ewt_amount,
                cheque_amount=cheque_amount,
                difference_amount=invoice_amount - cheque_amount,
            )
        )
    return rows


def _read_sheet1_columns(sheet):
    """Read Sheet1 where only B (cheque), D (invoice), and E (amount) are supplied."""
    rows = []
    cheque_date = sheet.cell(row=2, column=2).value
    # Streaming is essential here. Random-access ``sheet.cell`` calls on a
    # read-only worksheet repeatedly scan its XML and become extremely slow
    # for production files with more than a thousand rows.
    for index, values in enumerate(
        sheet.iter_rows(min_row=4, max_col=5, values_only=True), start=4
    ):
        cheque_no = values[1]
        invoice_no = values[3]
        source_amount = values[4]
        if cheque_no in (None, "") and invoice_no in (None, "") and source_amount in (None, ""):
            continue
        if cheque_no in (None, "") or invoice_no in (None, "") or source_amount in (None, ""):
            frappe.throw(_("Sheet1 row {0} must contain columns B, D, and E.").format(index))
        try:
            invoice_amount = _decimal(source_amount)
        except InvalidOperation:
            frappe.throw(_("Sheet1 row {0} contains an invalid amount in column E.").format(index))
        basis_amount = invoice_amount * Decimal("1.01")
        ewt_amount = basis_amount * Decimal("0.01")
        cheque_amount = basis_amount - ewt_amount
        difference_amount = invoice_amount - cheque_amount
        rows.append(
            frappe._dict(
                row_no=index,
                cheque_no=_clean_id(cheque_no),
                cheque_date=str(getdate(cheque_date)) if cheque_date else None,
                invoice_no=_clean_id(invoice_no),
                invoice_amount=invoice_amount,
                basis_amount=basis_amount,
                ewt_amount=ewt_amount,
                cheque_amount=cheque_amount,
                difference_amount=difference_amount,
            )
        )
    return rows


def _header(value):
    return " ".join(str(value or "").strip().lower().split())


def _clean_id(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\u00a0", " ").strip()


def _decimal(value):
    return Decimal(str(value or 0).replace(",", ""))


def _money(value):
    return _decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _display_date(value):
    value = getdate(value)
    return "{0}/{1}/{2}".format(value.month, value.day, value.year)


def _require_create_permission():
    if not frappe.has_permission("Journal Entry", "create"):
        frappe.throw(_("You do not have permission to create Journal Entries."), frappe.PermissionError)
